#!/usr/bin/env python3
"""OCI provisioned-resource inventory (IaaS & PaaS) -> single Excel workbook.

Standalone, read-only script that walks every subscribed region and every
accessible compartment of the tenancy and lists the resources that are
actually PROVISIONED - running, stopped / powered off, or being created.
Terminated, deleted and failed resources are left out, and so are
configuration objects (route tables, security lists, attachments, IAM
policies, users, alarms, ...). Each sheet holds only the key columns for
that resource type (name, status, shape, OCPUs, memory, storage, version).

Only OCI SDK ``list_*`` / ``get_*`` calls are used; nothing is ever created,
changed or deleted. Authentication reuses your OCI CLI / API-key setup
(``~/.oci/config``).

Usage:
    pip install oci openpyxl
    python3 oci_inventory_excel.py --profile TanencyB
    python3 oci_inventory_excel.py --regions us-ashburn-1,uk-london-1
    python3 oci_inventory_excel.py --compartment ocid1.compartment.oc1..xxxx
    python3 oci_inventory_excel.py --tags              # add tag columns

Output: ``<output-dir>/OCI_Inventory_<profile>_<timestamp>.xlsx`` with sheets
    Summary           resource count per resource type and region, with a
                      link to each sheet
    All Resources     every provisioned resource, one row each
    <Resource Type>   one sheet per resource type (Instance, Block Volume,
                      VCN, ATP, ADW, OAC, OKE Cluster, ...)
    Errors            API calls that failed (missing permissions, service
                      not available in a region, ...)
    Run Info          tenancy, profile, regions scanned, timestamp, totals
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import sys
import threading
import time
from collections import Counter, defaultdict
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
from datetime import datetime
from pathlib import Path

try:
    import oci
except ImportError:  # pragma: no cover
    sys.exit("The OCI Python SDK is required:  pip install oci openpyxl")
try:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font, PatternFill
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")

log = logging.getLogger("oci_inventory")

# Not provisioned: gone, going away, or never created successfully.
SKIP_STATES = {"TERMINATED", "TERMINATING", "DELETED", "DELETING", "FAILED"}
MAX_CELL_LEN = 32000  # stay under Excel's 32,767 character cell limit
MAX_ROWS = 1_048_575  # Excel row limit minus the header row

# The SDK default retries a failing call for up to 10 minutes. Some services
# have no endpoint in some regions, so fail fast instead: 3 attempts, 60s max.
RETRY_STRATEGY = oci.retry.RetryStrategyBuilder(
    max_attempts_check=True, max_attempts=3,
    total_elapsed_time_check=True, total_elapsed_time_seconds=60,
    retry_max_wait_between_calls_seconds=5,
    service_error_check=True,
).get_retry_strategy()
TIMEOUT = (5, 60)  # (connect, read) seconds

# ---------------------------------------------------------------------------
# Resource catalogue - provisioned resources only
# ---------------------------------------------------------------------------
# scope:
#   comp    list_x(compartment_id=<compartment>)
#   ad      list_x(compartment_id=<compartment>, availability_domain=<ad>) per AD
#   ns      list_x(namespace_name=<ns>, compartment_id=<compartment>)
# fields:  extra columns for the resource's sheet, as (column, attribute path).
#          A path is dotted for nested values ("shape_config.ocpus").
# subtype: (attribute, {value: type}) - split one list call into several
#          resource types by an attribute, e.g. Autonomous Databases into ATP/ADW.


def R(category, service, rtype, client, method, fields=(), scope="comp", kwargs=None, subtype=None):
    return {
        "category": category, "service": service, "type": rtype, "client": client, "method": method,
        "fields": list(fields), "scope": scope, "kwargs": kwargs or {}, "subtype": subtype,
    }


ADB_FIELDS = [
    ("Compute Model", "compute_model"), ("ECPU/OCPU Count", "compute_count"),
    ("Storage (TB)", "data_storage_size_in_tbs"), ("Auto Scaling", "is_auto_scaling_enabled"),
    ("DB Version", "db_version"), ("License", "license_model"), ("Always Free", "is_free_tier"),
    ("Dedicated", "is_dedicated"),
]

RESOURCES = [
    # ----------------------------- IaaS ----------------------------------
    R("IaaS", "Compute", "Instance", "core.ComputeClient", "list_instances",
      [("OCPUs", "shape_config.ocpus"), ("Memory (GB)", "shape_config.memory_in_gbs"),
       ("Fault Domain", "fault_domain")]),
    R("IaaS", "Compute", "Dedicated VM Host", "core.ComputeClient", "list_dedicated_vm_hosts",
      [("Total OCPUs", "total_ocpus"), ("Total Memory (GB)", "total_memory_in_gbs")]),
    R("IaaS", "Compute", "Instance Pool", "core.ComputeManagementClient", "list_instance_pools",
      [("Pool Size", "size")]),
    R("IaaS", "Compute", "Cluster Network", "core.ComputeManagementClient", "list_cluster_networks"),
    R("IaaS", "Block Storage", "Block Volume", "core.BlockstorageClient", "list_volumes",
      [("Size (GB)", "size_in_gbs"), ("VPUs per GB", "vpus_per_gb")]),
    R("IaaS", "Block Storage", "Boot Volume", "core.BlockstorageClient", "list_boot_volumes",
      [("Size (GB)", "size_in_gbs"), ("VPUs per GB", "vpus_per_gb")]),
    R("IaaS", "Block Storage", "Volume Group", "core.BlockstorageClient", "list_volume_groups",
      [("Size (GB)", "size_in_gbs")]),
    R("IaaS", "File Storage", "File System", "file_storage.FileStorageClient", "list_file_systems",
      [("Used (Bytes)", "metered_bytes")], scope="ad"),
    R("IaaS", "File Storage", "Mount Target", "file_storage.FileStorageClient", "list_mount_targets", scope="ad"),
    R("IaaS", "Object Storage", "Bucket", "object_storage.ObjectStorageClient", "list_buckets", scope="ns"),
    R("IaaS", "Networking", "VCN", "core.VirtualNetworkClient", "list_vcns",
      [("CIDR Blocks", "cidr_blocks"), ("DNS Label", "dns_label")]),
    R("IaaS", "Networking", "Subnet", "core.VirtualNetworkClient", "list_subnets",
      [("CIDR Block", "cidr_block"), ("Private Subnet", "prohibit_public_ip_on_vnic")]),
    R("IaaS", "Networking", "Internet Gateway", "core.VirtualNetworkClient", "list_internet_gateways"),
    R("IaaS", "Networking", "NAT Gateway", "core.VirtualNetworkClient", "list_nat_gateways",
      [("NAT IP", "nat_ip")]),
    R("IaaS", "Networking", "Service Gateway", "core.VirtualNetworkClient", "list_service_gateways"),
    R("IaaS", "Networking", "Local Peering Gateway", "core.VirtualNetworkClient", "list_local_peering_gateways",
      [("Peering Status", "peering_status")]),
    R("IaaS", "Networking", "DRG", "core.VirtualNetworkClient", "list_drgs"),
    R("IaaS", "Networking", "Remote Peering Connection", "core.VirtualNetworkClient",
      "list_remote_peering_connections", [("Peering Status", "peering_status")]),
    R("IaaS", "Networking", "Reserved Public IP", "core.VirtualNetworkClient", "list_public_ips",
      [("IP Address", "ip_address"), ("Assigned To", "assigned_entity_type")],
      kwargs={"scope": "REGION", "lifetime": "RESERVED"}),
    R("IaaS", "Networking", "Site-to-Site VPN", "core.VirtualNetworkClient", "list_ip_sec_connections",
      [("Static Routes", "static_routes")]),
    R("IaaS", "Networking", "FastConnect Virtual Circuit", "core.VirtualNetworkClient", "list_virtual_circuits",
      [("Bandwidth", "bandwidth_shape_name"), ("Type", "type"), ("Provider State", "provider_state")]),
    R("IaaS", "Networking", "Load Balancer", "load_balancer.LoadBalancerClient", "list_load_balancers",
      [("IP Addresses", "ip_addresses.ip_address"), ("Private", "is_private"),
       ("Min Bandwidth (Mbps)", "shape_details.minimum_bandwidth_in_mbps"),
       ("Max Bandwidth (Mbps)", "shape_details.maximum_bandwidth_in_mbps")]),
    R("IaaS", "Networking", "Network Load Balancer", "network_load_balancer.NetworkLoadBalancerClient",
      "list_network_load_balancers", [("IP Addresses", "ip_addresses.ip_address"), ("Private", "is_private")]),
    R("IaaS", "Networking", "DNS Zone", "dns.DnsClient", "list_zones", [("Zone Type", "zone_type")]),
    R("IaaS", "Networking", "Network Firewall", "network_firewall.NetworkFirewallClient", "list_network_firewalls",
      [("IP Address", "ipv4_address")]),
    R("IaaS", "Security", "Web App Firewall", "waf.WafClient", "list_web_app_firewalls",
      [("Backend Type", "backend_type")]),
    R("IaaS", "Security", "Vault", "key_management.KmsVaultClient", "list_vaults", [("Vault Type", "vault_type")]),
    R("IaaS", "Security", "Bastion", "bastion.BastionClient", "list_bastions", [("Bastion Type", "bastion_type")]),

    # ----------------------------- PaaS ----------------------------------
    R("PaaS", "Database", "DB System", "database.DatabaseClient", "list_db_systems",
      [("CPU Cores", "cpu_core_count"), ("Nodes", "node_count"), ("Storage (GB)", "data_storage_size_in_gbs"),
       ("Edition", "database_edition"), ("Version", "version"), ("License", "license_model")]),
    R("PaaS", "Database", "Autonomous Database", "database.DatabaseClient", "list_autonomous_databases",
      ADB_FIELDS, subtype=("db_workload", {
          "OLTP": "Autonomous Transaction Processing (ATP)",
          "DW": "Autonomous Data Warehouse (ADW)",
          "AJD": "Autonomous JSON Database (AJD)",
          "APEX": "Autonomous APEX Service",
          "LH": "Autonomous AI Lakehouse",
      })),
    R("PaaS", "Database", "Autonomous Container Database", "database.DatabaseClient",
      "list_autonomous_container_databases", [("Version", "db_version")]),
    R("PaaS", "Database", "Cloud Exadata Infrastructure", "database.DatabaseClient",
      "list_cloud_exadata_infrastructures",
      [("DB Servers", "compute_count"), ("Storage Servers", "storage_count")]),
    R("PaaS", "Database", "Exadata VM Cluster", "database.DatabaseClient", "list_cloud_vm_clusters",
      [("CPU Cores", "cpu_core_count"), ("Nodes", "node_count"), ("GI Version", "gi_version"),
       ("License", "license_model")]),
    R("PaaS", "Database", "Autonomous Exadata VM Cluster", "database.DatabaseClient",
      "list_cloud_autonomous_vm_clusters", [("CPU Cores", "cpu_core_count"), ("Nodes", "node_count")]),
    R("PaaS", "Database", "ExaDB-XS VM Cluster", "database.DatabaseClient", "list_exadb_vm_clusters",
      [("Nodes", "node_count"), ("ECPUs", "enabled_e_cpu_count"), ("License", "license_model")]),
    R("PaaS", "MySQL HeatWave", "MySQL DB System", "mysql.DbSystemClient", "list_db_systems",
      [("Version", "mysql_version"), ("HeatWave Attached", "is_heat_wave_cluster_attached")]),
    R("PaaS", "PostgreSQL", "PostgreSQL DB System", "psql.PostgresqlClient", "list_db_systems",
      [("Nodes", "instance_count"), ("OCPUs", "instance_ocpu_count"),
       ("Memory (GB)", "instance_memory_size_in_gbs"), ("Version", "db_version")]),
    R("PaaS", "NoSQL", "NoSQL Table", "nosql.NosqlClient", "list_tables",
      [("Max Reads/s", "table_limits.max_read_units"), ("Max Writes/s", "table_limits.max_write_units"),
       ("Max Storage (GB)", "table_limits.max_storage_in_g_bs"), ("Capacity Mode", "table_limits.capacity_mode")]),
    R("PaaS", "Cache", "Redis Cluster", "redis.RedisClusterClient", "list_redis_clusters",
      [("Nodes", "node_count"), ("Memory per Node (GB)", "node_memory_in_gbs"), ("Version", "software_version")]),
    R("PaaS", "Search", "OpenSearch Cluster", "opensearch.OpensearchClusterClient", "list_opensearch_clusters",
      [("Version", "software_version")]),
    R("PaaS", "GoldenGate", "GoldenGate Deployment", "golden_gate.GoldenGateClient", "list_deployments",
      [("Deployment Type", "deployment_type"), ("CPU Cores", "cpu_core_count"),
       ("Auto Scaling", "is_auto_scaling_enabled"), ("License", "license_model")]),
    R("PaaS", "Kubernetes (OKE)", "OKE Cluster", "container_engine.ContainerEngineClient", "list_clusters",
      [("Kubernetes Version", "kubernetes_version"), ("Cluster Type", "type")]),
    R("PaaS", "Kubernetes (OKE)", "OKE Node Pool", "container_engine.ContainerEngineClient", "list_node_pools",
      [("Nodes", "node_config_details.size"), ("OCPUs per Node", "node_shape_config.ocpus"),
       ("Memory per Node (GB)", "node_shape_config.memory_in_gbs"), ("Kubernetes Version", "kubernetes_version")]),
    R("PaaS", "Kubernetes (OKE)", "OKE Virtual Node Pool", "container_engine.ContainerEngineClient",
      "list_virtual_node_pools", [("Nodes", "size"), ("Kubernetes Version", "kubernetes_version")]),
    R("PaaS", "Container Instances", "Container Instance", "container_instances.ContainerInstanceClient",
      "list_container_instances",
      [("OCPUs", "shape_config.ocpus"), ("Memory (GB)", "shape_config.memory_in_gbs"),
       ("Containers", "container_count")]),
    R("PaaS", "Functions", "Functions Application", "functions.FunctionsManagementClient", "list_applications"),
    R("PaaS", "API Gateway", "API Gateway", "apigateway.GatewayClient", "list_gateways",
      [("Endpoint Type", "endpoint_type"), ("Hostname", "hostname")]),
    R("PaaS", "Integration", "Integration Instance", "integration.IntegrationInstanceClient",
      "list_integration_instances",
      [("Edition", "integration_instance_type"), ("Message Packs", "message_packs"), ("BYOL", "is_byol")]),
    R("PaaS", "Visual Builder", "Visual Builder Instance", "visual_builder.VbInstanceClient", "list_vb_instances",
      [("Nodes", "node_count")]),
    R("PaaS", "Analytics", "Oracle Analytics Cloud (OAC)", "analytics.AnalyticsClient", "list_analytics_instances",
      [("Edition", "feature_set"), ("Capacity Type", "capacity.capacity_type"),
       ("Capacity", "capacity.capacity_value"), ("License", "license_type")]),
    R("PaaS", "Digital Assistant", "Digital Assistant Instance", "oda.OdaClient", "list_oda_instances",
      [("Size", "shape_name")]),
    R("PaaS", "Data Science", "Notebook Session", "data_science.DataScienceClient", "list_notebook_sessions"),
    R("PaaS", "Data Science", "Model Deployment", "data_science.DataScienceClient", "list_model_deployments"),
    R("PaaS", "Data Flow", "Data Flow Application", "data_flow.DataFlowClient", "list_applications",
      [("Language", "language"), ("Spark Version", "spark_version"), ("Type", "type")]),
    R("PaaS", "Data Integration", "Data Integration Workspace", "data_integration.DataIntegrationClient",
      "list_workspaces"),
    R("PaaS", "Data Catalog", "Data Catalog", "data_catalog.DataCatalogClient", "list_catalogs"),
    R("PaaS", "Big Data", "Big Data Service Cluster", "bds.BdsClient", "list_bds_instances",
      [("Nodes", "number_of_nodes"), ("Version", "cluster_version"), ("High Availability", "is_high_availability")]),
    R("PaaS", "Generative AI", "GenAI Dedicated AI Cluster", "generative_ai.GenerativeAiClient",
      "list_dedicated_ai_clusters", [("Cluster Type", "type"), ("Units", "unit_count"), ("Unit Shape", "unit_shape")]),
    R("PaaS", "Generative AI", "GenAI Endpoint", "generative_ai.GenerativeAiClient", "list_endpoints"),
    R("PaaS", "Generative AI", "GenAI Agent", "generative_ai_agent.GenerativeAiAgentClient", "list_agents"),
    R("PaaS", "Generative AI", "GenAI Agent Endpoint", "generative_ai_agent.GenerativeAiAgentClient",
      "list_agent_endpoints"),
    R("PaaS", "Generative AI", "GenAI Knowledge Base", "generative_ai_agent.GenerativeAiAgentClient",
      "list_knowledge_bases"),
    R("PaaS", "Streaming", "Stream", "streaming.StreamAdminClient", "list_streams", [("Partitions", "partitions")]),
    R("PaaS", "Queue", "Queue", "queue.QueueAdminClient", "list_queues"),
    R("PaaS", "Blockchain", "Blockchain Platform", "blockchain.BlockchainPlatformClient",
      "list_blockchain_platforms", [("Platform Role", "platform_role")]),
    R("PaaS", "APM", "APM Domain", "apm_control_plane.ApmDomainClient", "list_apm_domains",
      [("Always Free", "is_free_tier")]),
]

BASE_COLUMNS = ["Region", "Compartment", "Service", "Resource Type", "Name", "Status", "Shape",
                "Availability Domain", "Created", "OCID"]
TAG_COLUMNS = ["Freeform Tags", "Defined Tags"]

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


class Inventory:
    """Thread-safe store for rows and errors, plus a per-thread client cache."""

    def __init__(self, config: dict, timeout: tuple[int, int]):
        self.config = config
        self.timeout = timeout
        self.rows: dict[tuple[str, str, str], list[dict]] = defaultdict(list)
        self.fields: dict[tuple[str, str, str], list[str]] = {}
        self.errors: list[dict] = []
        self.unreachable: set[tuple[str, str]] = set()  # (client, region) with no reachable endpoint
        self.probed: set[tuple[str, str]] = set()       # (client, region) whose first call has returned
        self._probe_locks: dict[tuple[str, str], threading.Lock] = defaultdict(threading.Lock)
        self._lock = threading.Lock()
        self._local = threading.local()

    def client(self, path: str, region: str):
        cache = getattr(self._local, "clients", None)
        if cache is None:
            cache = self._local.clients = {}
        key = (path, region)
        if key not in cache:
            module_name, class_name = path.split(".")
            cls = getattr(getattr(oci, module_name), class_name)
            cfg = dict(self.config, region=region)
            cache[key] = cls(cfg, retry_strategy=RETRY_STRATEGY, timeout=self.timeout)
        return cache[key]

    def probe_lock(self, key) -> threading.Lock:
        with self._lock:
            return self._probe_locks[key]

    def add(self, key, rows, fields):
        with self._lock:
            self.rows[key].extend(rows)
            self.fields[key] = fields

    def error(self, region, spec_type, compartment, exc):
        status = getattr(exc, "status", "")
        code = getattr(exc, "code", type(exc).__name__)
        message = getattr(exc, "message", None) or str(exc)
        with self._lock:
            self.errors.append({
                "Region": region, "Resource Type": spec_type, "Compartment": compartment,
                "Status": status, "Code": code, "Message": message[:500],
            })


def list_all(client, method: str, **kwargs) -> list:
    fn = getattr(client, method)
    return oci.pagination.list_call_get_all_results(fn, **kwargs).data


def pick(data, path: str):
    """Read a dotted attribute path; lists are walked and joined ("a, b")."""
    values = [data]
    for part in path.split("."):
        nxt = []
        for v in values:
            if isinstance(v, list):
                nxt += [i.get(part) if isinstance(i, dict) else None for i in v]
            elif isinstance(v, dict):
                nxt.append(v.get(part))
        values = nxt
    flat = []
    for v in values:
        flat += v if isinstance(v, list) else [v]
    flat = [v for v in flat if v not in (None, "", {})]
    if not flat:
        return ""
    if len(flat) == 1 and not isinstance(flat[0], (dict, list)):
        return flat[0]
    return ", ".join(json.dumps(v, default=str) if isinstance(v, (dict, list)) else str(v) for v in flat)


def to_json(value) -> str:
    return json.dumps(value, default=str, sort_keys=True) if value else ""


def build_row(item, spec, region, comp, with_tags: bool) -> dict:
    data = oci.util.to_dict(item)
    if not isinstance(data, dict):
        data = {}
    shape = (data.get("shape") or data.get("shape_name") or data.get("node_shape")
             or data.get("dedicated_vm_host_shape") or data.get("storage_tier") or "")
    row = {
        "Region": region,
        "Compartment": comp["path"],
        "Service": spec["service"],
        "Resource Type": spec["type"],
        "Name": data.get("display_name") or data.get("name") or "",
        # A listed resource with no lifecycle state (e.g. a bucket) simply exists.
        "Status": data.get("lifecycle_state") or "ACTIVE",
        "Shape": shape if isinstance(shape, str) else "",
        "Availability Domain": data.get("availability_domain") or "",
        "Created": str(data.get("time_created") or ""),
        "OCID": data.get("id") or "",
        "_compartment_id": data.get("compartment_id") or "",
        "_data": data,
    }
    for column, path in spec["fields"]:
        row[column] = pick(data, path)
    if with_tags:
        row["Freeform Tags"] = to_json(data.get("freeform_tags"))
        row["Defined Tags"] = to_json(data.get("defined_tags"))
    return row


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")


# ---------------------------------------------------------------------------
# Discovery
# ---------------------------------------------------------------------------


def get_compartments(identity, tenancy_id: str, only: str | None) -> list[dict]:
    tenancy = identity.get_compartment(tenancy_id).data
    items = list_all(identity, "list_compartments", compartment_id=tenancy_id,
                     compartment_id_in_subtree=True, access_level="ACCESSIBLE", lifecycle_state="ACTIVE")
    by_id = {tenancy_id: {"id": tenancy_id, "name": f"{tenancy.name} (root)", "parent": None}}
    for c in items:
        by_id[c.id] = {"id": c.id, "name": c.name, "parent": c.compartment_id}

    def path(cid):
        parts, seen = [], set()
        while cid in by_id and cid not in seen:
            seen.add(cid)
            parts.append(by_id[cid]["name"])
            cid = by_id[cid]["parent"]
        return " / ".join(reversed(parts))

    comps = [{"id": c["id"], "name": c["name"], "path": path(c["id"])} for c in by_id.values()]
    if only:
        # The selected compartment plus all of its sub-compartments.
        wanted = {only}
        changed = True
        while changed:
            changed = False
            for c in by_id.values():
                if c["parent"] in wanted and c["id"] not in wanted:
                    wanted.add(c["id"])
                    changed = True
        comps = [c for c in comps if c["id"] in wanted]
        if not comps:
            raise SystemExit(f"Compartment {only} not found or not accessible.")
    return sorted(comps, key=lambda c: c["path"])


def get_regions(identity, tenancy_id: str, only: list[str] | None) -> tuple[list[str], str]:
    subs = identity.list_region_subscriptions(tenancy_id).data
    home = next((s.region_name for s in subs if s.is_home_region), subs[0].region_name)
    regions = sorted(s.region_name for s in subs if s.status == "READY")
    if only:
        unknown = set(only) - set(regions)
        if unknown:
            log.warning("Not subscribed / unknown regions ignored: %s", ", ".join(sorted(unknown)))
        regions = [r for r in regions if r in only]
    return regions, home


# ---------------------------------------------------------------------------
# Collection
# ---------------------------------------------------------------------------


def collect_task(inv: Inventory, spec, region, comp, ctx, args) -> int:
    """List one resource type in one compartment of one region."""
    key = (spec["client"], region)

    def fetch():
        client = inv.client(spec["client"], region)
        kwargs = dict(spec["kwargs"])
        if spec["scope"] == "ad":
            items = []
            for ad in ctx["ads"].get(region, []):
                items += list_all(client, spec["method"], compartment_id=comp["id"], availability_domain=ad, **kwargs)
            return items
        if spec["scope"] == "ns":
            return list_all(client, spec["method"], namespace_name=ctx["namespace"], compartment_id=comp["id"],
                            **kwargs)
        return list_all(client, spec["method"], compartment_id=comp["id"], **kwargs)

    try:
        if key in inv.unreachable:
            return 0  # endpoint already found unreachable in this region; reported once
        if key in inv.probed:
            items = fetch()
        else:
            # The first call to a service in a region runs alone, so an
            # unreachable endpoint costs one failed call, not one per compartment.
            with inv.probe_lock(key):
                if key in inv.unreachable:
                    return 0
                try:
                    items = fetch()
                finally:
                    inv.probed.add(key)
    except oci.exceptions.ServiceError as exc:
        inv.error(region, spec["type"], comp["path"], exc)
        return 0
    except Exception as exc:  # noqa: BLE001 - one failure must never stop the run
        # Connection / DNS / timeout errors: the service is not reachable in this
        # region, so skip it for every other compartment too.
        if key not in inv.unreachable:
            inv.unreachable.add(key)
            inv.error(region, spec["type"], "(all compartments)",
                      RuntimeError(f"Service endpoint not reachable in {region}; skipped: {exc}"))
        return 0

    fields = [c for c, _ in spec["fields"]]
    grouped = defaultdict(list)
    for item in items:
        row = build_row(item, spec, region, comp, args.tags)
        # Only resources that live in this compartment (some APIs also return
        # objects from elsewhere), and only ones that are actually provisioned.
        if row["_compartment_id"] and row["_compartment_id"] != comp["id"]:
            continue
        if not args.include_deleted and str(row["Status"]).upper() in SKIP_STATES:
            continue
        if spec["subtype"]:
            attr, names = spec["subtype"]
            value = row["_data"].get(attr) or ""
            row["Resource Type"] = names.get(value, f"{spec['type']} ({value or 'unknown'})")
        grouped[row["Resource Type"]].append(row)

    for rtype, rows in grouped.items():
        inv.add((spec["category"], spec["service"], rtype), rows, fields)
    return sum(len(r) for r in grouped.values())


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------


def cell_value(value):
    """Make a value safe for an Excel cell."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > MAX_CELL_LEN:
            value = value[:MAX_CELL_LEN] + "...[truncated]"
    return value


class SheetNamer:
    """Excel sheet names: max 31 chars, no []:*?/\\, unique (case-insensitive)."""

    def __init__(self):
        self.used = set()

    def __call__(self, text: str) -> str:
        base = re.sub(r"[\[\]:*?/\\]", "-", text).strip("' ")[:31] or "Sheet"
        name, n = base, 2
        while name.lower() in self.used:
            suffix = f" ({n})"
            name, n = base[:31 - len(suffix)] + suffix, n + 1
        self.used.add(name.lower())
        return name


# Resource types whose names exceed Excel's 31-character sheet-name limit.
SHEET_ALIASES = {
    "Autonomous Transaction Processing (ATP)": "Autonomous Transaction (ATP)",
}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="C74634")
LINK_FONT = Font(color="0563C1", underline="single")


def add_sheet(wb, title: str, rows: list[dict], columns: list[str], links: dict | None = None) -> None:
    """Append a formatted sheet: bold coloured header, frozen top row, filter, sized columns."""
    ws = wb.create_sheet(title)
    if len(rows) > MAX_ROWS:
        log.warning("Sheet '%s' has %d rows; only the first %d fit in Excel.", title, len(rows), MAX_ROWS)
        rows = rows[:MAX_ROWS]
    for idx, col in enumerate(columns, start=1):
        sample = [len(str(r.get(col, ""))) for r in rows[:200]]
        width = min(max([len(col)] + sample) + 2, 60)
        ws.column_dimensions[_col_letter(idx)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{_col_letter(len(columns))}{len(rows) + 1}"

    header = []
    for col in columns:
        cell = WriteOnlyCell(ws, value=col)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        header.append(cell)
    ws.append(header)
    for row in rows:
        values = []
        for col in columns:
            value = cell_value(row.get(col, ""))
            if links and col in links and row.get(col) in links[col]:
                target = links[col][row[col]].replace("'", "''")
                cell = WriteOnlyCell(ws, value=f'=HYPERLINK("#\'{target}\'!A1","{value}")')
                cell.font = LINK_FONT
                value = cell
            values.append(value)
        ws.append(values)


def _col_letter(n: int) -> str:
    letters = ""
    while n:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def write_workbook(inv: Inventory, path: Path, meta: dict, with_tags: bool) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    namer = SheetNamer()
    for reserved in ("Summary", "All Resources", "Errors", "Run Info"):
        namer(reserved)
    tags = TAG_COLUMNS if with_tags else []

    all_rows, summary, type_sheets, sheet_for_type = [], Counter(), [], {}
    for key in sorted(k for k, v in inv.rows.items() if v):
        category, service, rtype = key
        rows = sorted(inv.rows[key], key=lambda r: (r["Region"], r["Compartment"], str(r["Name"]).lower()))
        fields = inv.fields[key]
        for r in rows:
            summary[(category, service, rtype, r["Region"])] += 1
            r["Key Details"] = "; ".join(f"{f}: {r[f]}" for f in fields if r.get(f) not in ("", None))
        all_rows.extend(rows)
        sheet = namer(SHEET_ALIASES.get(rtype, rtype))
        sheet_for_type[rtype] = sheet
        # Sheet column order: identity, then the type's key fields, then the rest.
        columns = BASE_COLUMNS[:7] + fields + BASE_COLUMNS[7:] + tags
        type_sheets.append((sheet, rows, columns))

    summary_rows = [{"Category": k[0], "Service": k[1], "Resource Type": k[2], "Region": k[3], "Count": v,
                     "Sheet": sheet_for_type[k[2]]} for k, v in sorted(summary.items())]
    info_rows = [{"Setting": k, "Value": v} for k, v in meta.items()]
    info_rows.append({"Setting": "Total Provisioned Resources", "Value": len(all_rows)})
    info_rows.append({"Setting": "Errors / Skipped Calls", "Value": len(inv.errors)})

    wb = Workbook(write_only=True)
    add_sheet(wb, "Summary", summary_rows, ["Category", "Service", "Resource Type", "Region", "Count", "Sheet"],
              links={"Sheet": {s: s for s in sheet_for_type.values()}})
    add_sheet(wb, "All Resources", all_rows, BASE_COLUMNS[:7] + ["Key Details"] + BASE_COLUMNS[7:] + tags)
    for sheet, rows, columns in type_sheets:
        add_sheet(wb, sheet, rows, columns)
    add_sheet(wb, "Errors", inv.errors, ["Region", "Resource Type", "Compartment", "Status", "Code", "Message"])
    add_sheet(wb, "Run Info", info_rows, ["Setting", "Value"])
    wb.save(path)
    log.info("Total provisioned resources: %d   (%d resource types)", len(all_rows), len(type_sheets))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Export provisioned OCI IaaS & PaaS resources for all regions to one Excel workbook (read-only).")
    p.add_argument("--config-file", default=os.getenv("OCI_CLI_CONFIG_FILE", oci.config.DEFAULT_LOCATION),
                   help="OCI config file (default: ~/.oci/config)")
    p.add_argument("--profile", default=os.getenv("OCI_CLI_PROFILE", "DEFAULT"), help="OCI config profile")
    p.add_argument("--regions", help="Comma-separated regions to scan (default: all subscribed regions)")
    p.add_argument("--compartment", help="Only scan this compartment OCID and its sub-compartments")
    p.add_argument("--output-dir", default="./output", help="Output directory (default: ./output)")
    p.add_argument("--workers", type=int, default=10, help="Parallel API workers (default: 10)")
    p.add_argument("--tags", action="store_true", help="Add Freeform / Defined tag columns")
    p.add_argument("--include-deleted", action="store_true",
                   help="Also export TERMINATED / DELETED / FAILED resources")
    p.add_argument("-v", "--verbose", action="store_true", help="Debug logging")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(asctime)s %(levelname)-7s %(message)s", datefmt="%H:%M:%S")
    if not args.verbose:
        logging.getLogger("oci").setLevel(logging.WARNING)

    config = oci.config.from_file(os.path.expanduser(args.config_file), args.profile)
    oci.config.validate_config(config)
    tenancy_id = config["tenancy"]
    inv = Inventory(config, timeout=TIMEOUT)

    identity = inv.client("identity.IdentityClient", config["region"])
    regions, home = get_regions(identity, tenancy_id, args.regions.split(",") if args.regions else None)
    identity = inv.client("identity.IdentityClient", home)
    comps = get_compartments(identity, tenancy_id, args.compartment)
    log.info("Tenancy: %s | Home region: %s", tenancy_id, home)
    log.info("Regions (%d): %s", len(regions), ", ".join(regions))
    log.info("Compartments: %d", len(comps))

    ctx = {"ads": {}, "namespace": None}
    try:
        ctx["namespace"] = inv.client("object_storage.ObjectStorageClient", home).get_namespace().data
    except Exception as exc:  # noqa: BLE001
        inv.error(home, "Object Storage Namespace", "", exc)
    for region in regions:
        try:
            ads = inv.client("identity.IdentityClient", region).list_availability_domains(tenancy_id).data
            ctx["ads"][region] = [ad.name for ad in ads]
        except Exception as exc:  # noqa: BLE001
            inv.error(region, "Availability Domains", "", exc)

    tasks = [(spec, region, comp) for spec in RESOURCES if spec["scope"] != "ns" or ctx["namespace"]
             for region in regions for comp in comps]

    log.info("Running %d list operations with %d workers ...", len(tasks), args.workers)
    started = last_log = time.time()
    done = 0
    with ThreadPoolExecutor(max_workers=max(1, args.workers)) as pool:
        futures = {pool.submit(collect_task, inv, s, r, c, ctx, args): (s, r, c) for s, r, c in tasks}
        pending = set(futures)
        while pending:
            finished, pending = wait(pending, timeout=30, return_when=FIRST_COMPLETED)
            for fut in finished:
                done += 1
                spec, region, comp = futures[fut]
                found = fut.result()
                if found:
                    log.debug("%-16s %-35s %-40s %d", region, spec["type"], comp["path"], found)
                if done % 250 == 0 or done == len(tasks):
                    log.info("Progress: %d/%d operations (%.0fs)", done, len(tasks), time.time() - started)
                    last_log = time.time()
            if pending and time.time() - last_log >= 30:
                running = [futures[f] for f in pending if f.running()][:5]
                log.info("Progress: %d/%d operations (%.0fs) - waiting on: %s", done, len(tasks),
                         time.time() - started, "; ".join(f"{s['type']} @ {r}" for s, r, _ in running))
                last_log = time.time()

    started_at = datetime.now()
    path = Path(args.output_dir) / f"OCI_Inventory_{safe_name(args.profile)}_{started_at:%Y%m%d_%H%M%S}.xlsx"
    meta = {
        "Tenancy": tenancy_id, "Profile": args.profile, "Home Region": home,
        "Regions Scanned": ", ".join(regions), "Compartments Scanned": len(comps),
        "Generated": f"{started_at:%Y-%m-%d %H:%M:%S}",
        "Included": "Provisioned resources only (running, stopped / powered off, provisioning)",
    }
    write_workbook(inv, path, meta, args.tags)
    log.info("Errors / skipped calls: %d (see the 'Errors' sheet)", len(inv.errors))
    log.info("Excel workbook written to: %s", path.resolve())


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit("Interrupted.")
